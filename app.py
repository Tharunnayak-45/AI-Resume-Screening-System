import os
import re
import json
from datetime import datetime

import mysql.connector
from mysql.connector import Error

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash
)

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from werkzeug.utils import secure_filename

from PyPDF2 import PdfReader
from docx import Document


# ============================================================
# FLASK APPLICATION
# ============================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "ai_resume_screening_secret_key"
)

app.config["UPLOAD_FOLDER"] = os.path.join(
    app.root_path,
    "uploads"
)

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx"
}


# ============================================================
# MYSQL CONFIGURATION
# ============================================================

DB_CONFIG = {
    "host": os.environ.get("MYSQLHOST", "localhost"),
    "port": int(
        os.environ.get(
            "MYSQLPORT",
            3306
        )
    ),
    "user": os.environ.get(
        "MYSQLUSER",
        "root"
    ),
    "password": os.environ.get(
        "MYSQLPASSWORD",
        "root"
    ),
    "database": os.environ.get(
        "MYSQLDATABASE",
        "ai_resume_db"
    ),
    "charset": "utf8mb4"
}


# ============================================================
# SKILL ALIASES
# ============================================================

SKILL_ALIASES = {

    "Java": [
        "java",
        "core java"
    ],

    "Python": [
        "python"
    ],

    "C": [
        "c",
        "c programming"
    ],

    "C++": [
        "c++"
    ],

    "C#": [
        "c#"
    ],

    "HTML": [
        "html",
        "html5"
    ],

    "CSS": [
        "css",
        "css3"
    ],

    "JavaScript": [
        "javascript",
        "js"
    ],

    "React": [
        "react",
        "react.js",
        "reactjs"
    ],

    "Angular": [
        "angular"
    ],

    "Node.js": [
        "node.js",
        "nodejs",
        "node js"
    ],

    "Flask": [
        "flask"
    ],

    "Django": [
        "django"
    ],

    "Spring": [
        "spring",
        "spring boot"
    ],

    "JSP": [
        "jsp"
    ],

    "Servlets": [
        "servlet",
        "servlets"
    ],

    "SQL": [
        "sql"
    ],

    "MySQL": [
        "mysql"
    ],

    "PostgreSQL": [
        "postgresql",
        "postgres"
    ],

    "Oracle": [
        "oracle"
    ],

    "MongoDB": [
        "mongodb",
        "mongo db"
    ],

    "JDBC": [
        "jdbc"
    ],

    "Git": [
        "git"
    ],

    "GitHub": [
        "github",
        "git hub"
    ],

    "VS Code": [
        "vs code",
        "visual studio code"
    ],

    "Apache Tomcat": [
        "apache tomcat",
        "tomcat"
    ],

    "Machine Learning": [
        "machine learning",
        "machine-learning"
    ],

    "Deep Learning": [
        "deep learning",
        "deep-learning"
    ],

    "Artificial Intelligence": [
        "artificial intelligence",
        "artificial-intelligence",
        "ai"
    ],

    "NLP": [
        "natural language processing",
        "nlp"
    ],

    "Pandas": [
        "pandas"
    ],

    "NumPy": [
        "numpy"
    ],

    "Scikit-learn": [
        "scikit-learn",
        "sklearn"
    ],

    "TensorFlow": [
        "tensorflow"
    ],

    "Data Structures": [
        "data structures",
        "data structure"
    ],

    "Algorithms": [
        "algorithms",
        "algorithm"
    ],

    "OOP": [
        "object oriented programming",
        "object-oriented programming",
        "oop"
    ],

    "DBMS": [
        "database management system",
        "dbms"
    ],

    "Operating Systems": [
        "operating systems",
        "operating system",
        "os"
    ],

    "Computer Networks": [
        "computer networks",
        "computer network",
        "cn"
    ],

    "Software Testing": [
        "software testing",
        "testing",
        "test cases",
        "functional testing",
        "unit testing"
    ],

    "Salesforce": [
        "salesforce"
    ],

    "Linux": [
        "linux"
    ]
}


# ============================================================
# ACTION VERBS
# ============================================================

ACTION_VERBS = [
    "developed",
    "develop",
    "designed",
    "design",
    "implemented",
    "implement",
    "created",
    "create",
    "built",
    "build",
    "developing",
    "managed",
    "manage",
    "tested",
    "test",
    "analyzed",
    "analyze",
    "improved",
    "improve",
    "optimized",
    "optimize",
    "integrated",
    "integrate",
    "automated",
    "automate",
    "deployed",
    "deploy",
    "configured",
    "configure",
    "maintained",
    "maintain",
    "debugged",
    "debug",
    "validated",
    "validate"
]


# ============================================================
# DATABASE CONNECTION
# ============================================================

def get_db_connection():

    connection = None

    try:

        connection = mysql.connector.connect(
            host=DB_CONFIG["host"],
            port=DB_CONFIG["port"],
            user=DB_CONFIG["user"],
            password=DB_CONFIG["password"],
            database=DB_CONFIG["database"],
            charset=DB_CONFIG["charset"]
        )

        if connection.is_connected():

            return connection

    except Error as e:

        print(
            "Database connection error:",
            e
        )

    return None


# ============================================================
# INITIALIZE DATABASE
# ============================================================

def initialize_database():

    connection = None
    cursor = None

    try:

        # ----------------------------------------------------
        # Connect to MySQL server
        # ----------------------------------------------------

        connection = mysql.connector.connect(
            host=DB_CONFIG["host"],
            port=DB_CONFIG["port"],
            user=DB_CONFIG["user"],
            password=DB_CONFIG["password"]
        )

        cursor = connection.cursor()

        cursor.execute(
            """
            CREATE DATABASE IF NOT EXISTS ai_resume_db
            CHARACTER SET utf8mb4
            COLLATE utf8mb4_unicode_ci
            """
        )

        connection.commit()

        cursor.close()
        connection.close()

        connection = None
        cursor = None

        # ----------------------------------------------------
        # Connect to application database
        # ----------------------------------------------------

        connection = get_db_connection()

        if connection is None:

            print(
                "Could not connect to ai_resume_db."
            )

            return False

        cursor = connection.cursor()

        # ----------------------------------------------------
        # USERS TABLE
        # ----------------------------------------------------

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                full_name VARCHAR(100) NOT NULL,
                email VARCHAR(150) UNIQUE NOT NULL,
                password VARCHAR(255) NOT NULL,
                role VARCHAR(30) DEFAULT 'user',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
            """
        )

        # ----------------------------------------------------
        # RESUMES TABLE
        # ----------------------------------------------------

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS resumes (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                filename VARCHAR(255) NOT NULL,
                extracted_text LONGTEXT,
                skills TEXT,
                uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                FOREIGN KEY (user_id)
                REFERENCES users(id)
                ON DELETE CASCADE
            )
            """
        )

        # ----------------------------------------------------
        # JOB DESCRIPTIONS TABLE
        # ----------------------------------------------------

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS job_descriptions (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT,
                job_title VARCHAR(150),
                company VARCHAR(150),
                location VARCHAR(150),
                description TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
            """
        )

        # ----------------------------------------------------
        # IMPORTANT
        #
        # The jobs table already exists in your database.
        #
        # This application does NOT create, alter, or modify
        # the existing jobs table.
        # ----------------------------------------------------

        # ----------------------------------------------------
        # ATS SCORES TABLE
        # ----------------------------------------------------

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS ats_scores (
                id INT AUTO_INCREMENT PRIMARY KEY,
                resume_id INT NOT NULL,
                job_description_id INT,

                overall_score DECIMAL(5,2),
                skills_score DECIMAL(5,2),
                section_score DECIMAL(5,2),
                contact_score DECIMAL(5,2),
                length_score DECIMAL(5,2),
                action_score DECIMAL(5,2),
                quantified_score DECIMAL(5,2),
                keyword_score DECIMAL(5,2),
                keyword_match_percentage DECIMAL(5,2),

                matched_keywords TEXT,
                missing_keywords TEXT,
                detected_skills TEXT,
                action_verbs TEXT,
                quantified_results TEXT,
                suggestions TEXT,

                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                FOREIGN KEY (resume_id)
                REFERENCES resumes(id)
                ON DELETE CASCADE
            )
            """
        )

        # ----------------------------------------------------
        # ANALYSIS RESULTS TABLE
        # ----------------------------------------------------

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS analysis_results (
                id INT AUTO_INCREMENT PRIMARY KEY,
                resume_id INT NOT NULL,
                resume_score DECIMAL(5,2),
                text_length INT,
                skills_count INT,
                detected_skills TEXT,
                suggestions TEXT,
                analyzed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                FOREIGN KEY (resume_id)
                REFERENCES resumes(id)
                ON DELETE CASCADE
            )
            """
        )

        connection.commit()

        print(
            "Database initialized successfully."
        )

        return True

    except Error as e:

        print(
            "Database initialization error:",
            e
        )

        return False

    finally:

        if cursor:

            try:
                cursor.close()
            except Exception:
                pass

        if connection:

            try:

                if connection.is_connected():
                    connection.close()

            except Exception:
                pass


# ============================================================
# FILE VALIDATION
# ============================================================

def allowed_file(filename):

    return (
        "." in filename
        and filename.rsplit(
            ".",
            1
        )[1].lower() in ALLOWED_EXTENSIONS
    )


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):

    if not text:

        return ""

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"\r\n|\r",
        "\n",
        text
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):

    text_parts = []

    try:

        reader = PdfReader(
            file_path
        )

        for page in reader.pages:

            try:

                page_text = page.extract_text()

                if page_text:

                    text_parts.append(
                        page_text
                    )

            except Exception as e:

                print(
                    "PDF page extraction error:",
                    e
                )

    except Exception as e:

        print(
            "PDF extraction error:",
            e
        )

    return normalize_text(
        "\n".join(text_parts)
    )


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_path):

    text_parts = []

    try:

        document = Document(
            file_path
        )

        # ----------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text_parts.append(
                    paragraph.text
                )

        # ----------------------------------------------------
        # Tables
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:

                        row_text.append(
                            cell_text
                        )

                if row_text:

                    text_parts.append(
                        " | ".join(row_text)
                    )

    except Exception as e:

        print(
            "DOCX extraction error:",
            e
        )

    return normalize_text(
        "\n".join(text_parts)
    )


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_pdf_text(
            file_path
        )

    if extension == ".docx":

        return extract_docx_text(
            file_path
        )

    return ""


# ============================================================
# SKILL DETECTION
# ============================================================

def detect_skills(text):

    if not text:

        return []

    lower_text = text.lower()

    found_skills = []

    for skill, aliases in SKILL_ALIASES.items():

        for alias in aliases:

            alias = alias.lower().strip()

            if not alias:

                continue

            # ------------------------------------------------
            # Special C handling
            # ------------------------------------------------

            if skill == "C":

                if re.search(
                    r"\bc(?:\s+programming)?\b",
                    lower_text
                ):

                    found_skills.append(
                        skill
                    )

                    break

            else:

                pattern = (
                    r"(?<![a-z0-9])"
                    + re.escape(alias)
                    + r"(?![a-z0-9])"
                )

                if re.search(
                    pattern,
                    lower_text
                ):

                    found_skills.append(
                        skill
                    )

                    break

    return sorted(
        set(found_skills)
    )


# ============================================================
# ACTION VERB DETECTION
# ============================================================

def detect_action_verbs(text):

    if not text:

        return []

    lower_text = text.lower()

    found = []

    for verb in ACTION_VERBS:

        pattern = (
            r"\b"
            + re.escape(verb)
            + r"\b"
        )

        if re.search(
            pattern,
            lower_text
        ):

            found.append(
                verb
            )

    return sorted(
        set(found)
    )


# ============================================================
# QUANTIFIED RESULTS
# ============================================================

def detect_quantified_results(text):

    if not text:

        return []

    patterns = [

        r"\b\d+(?:\.\d+)?\s*%",

        r"\b\d+(?:\.\d+)?\s*"
        r"(?:users|clients|customers|projects|pages|records|modules|days|months|years)\b",

        r"\b\d+(?:\.\d+)?\s*(?:k|m)\+?\b",

        r"\b\d+\+\b"
    ]

    results = []

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        for match in matches:

            clean_match = match.strip()

            if clean_match not in results:

                results.append(
                    clean_match
                )

    return results


# ============================================================
# CONTACT SCORE
# MAX = 6
# ============================================================

def calculate_contact_score(text):

    if not text:

        return 0

    score = 0

    lower_text = text.lower()

    # --------------------------------------------------------
    # Email
    # --------------------------------------------------------

    if re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    ):

        score += 2

    # --------------------------------------------------------
    # Phone
    # --------------------------------------------------------

    if re.search(
        r"(?:\+91[\s-]?)?[6-9]\d{9}",
        text
    ):

        score += 2

    # --------------------------------------------------------
    # LinkedIn
    # --------------------------------------------------------

    if "linkedin.com" in lower_text:

        score += 1

    # --------------------------------------------------------
    # GitHub
    # --------------------------------------------------------

    if "github.com" in lower_text:

        score += 1

    return min(
        score,
        6
    )


# ============================================================
# SECTION SCORE
# MAX = 22
# ============================================================

def calculate_section_score(text):

    if not text:

        return 0

    lower_text = text.lower()

    sections = {

        "summary": [
            "professional summary",
            "summary",
            "profile",
            "objective"
        ],

        "skills": [
            "technical skills",
            "skills"
        ],

        "education": [
            "education",
            "academic"
        ],

        "experience": [
            "experience",
            "work experience",
            "internship",
            "internships"
        ],

        "projects": [
            "projects",
            "project"
        ],

        "certifications": [
            "certifications",
            "certification",
            "certificates"
        ],

        "achievements": [
            "achievements",
            "achievement",
            "accomplishments"
        ]
    }

    points = 0

    for section_name, keywords in sections.items():

        found = False

        for keyword in keywords:

            if keyword in lower_text:

                found = True

                break

        if found:

            if section_name in [
                "summary",
                "skills",
                "education",
                "experience",
                "projects"
            ]:

                points += 4

            else:

                points += 2

    return min(
        points,
        22
    )


# ============================================================
# LENGTH SCORE
# MAX = 10
# ============================================================

def calculate_length_score(text):

    if not text:

        return 0

    length = len(
        text.strip()
    )

    if length >= 3000:

        return 10

    if length >= 2000:

        return 8

    if length >= 1200:

        return 6

    if length >= 700:

        return 4

    if length >= 400:

        return 2

    return 0


# ============================================================
# SKILL SCORE
# MAX = 20
# ============================================================

def calculate_skill_score(skills):

    count = len(skills)

    if count >= 15:

        return 20

    if count >= 12:

        return 17

    if count >= 9:

        return 14

    if count >= 6:

        return 10

    if count >= 3:

        return 6

    if count >= 1:

        return 3

    return 0


# ============================================================
# ACTION SCORE
# MAX = 10
# ============================================================

def calculate_action_score(action_verbs):

    count = len(action_verbs)

    if count >= 10:

        return 10

    if count >= 8:

        return 8

    if count >= 5:

        return 6

    if count >= 3:

        return 4

    if count >= 1:

        return 2

    return 0


# ============================================================
# QUANTIFIED SCORE
# MAX = 7
# ============================================================

def calculate_quantified_score(results):

    count = len(results)

    if count >= 5:

        return 7

    if count >= 4:

        return 6

    if count >= 3:

        return 5

    if count >= 2:

        return 3

    if count >= 1:

        return 1

    return 0


# ============================================================
# JOB KEYWORDS
# ============================================================

def extract_job_keywords(job_description):

    if not job_description:

        return []

    return detect_skills(
        job_description
    )


# ============================================================
# KEYWORD MATCH
# MAX = 25
# ============================================================

def calculate_keyword_match(
    resume_text,
    job_description
):

    if not job_description:

        return (
            0,
            0,
            [],
            []
        )

    resume_lower = resume_text.lower()

    job_keywords = extract_job_keywords(
        job_description
    )

    if not job_keywords:

        return (
            0,
            0,
            [],
            []
        )

    matched_keywords = []

    missing_keywords = []

    for keyword in job_keywords:

        if keyword.lower() in resume_lower:

            matched_keywords.append(
                keyword
            )

        else:

            missing_keywords.append(
                keyword
            )

    match_percentage = (
        len(matched_keywords)
        / len(job_keywords)
    ) * 100

    keyword_score = (
        match_percentage
        / 100
    ) * 25

    return (
        round(
            keyword_score,
            2
        ),
        round(
            match_percentage,
            2
        ),
        sorted(
            set(matched_keywords)
        ),
        sorted(
            set(missing_keywords)
        )
    )


# ============================================================
# SUGGESTIONS
# ============================================================

def generate_suggestions(
    text,
    skills,
    action_verbs,
    quantified_results,
    matched_keywords,
    missing_keywords
):

    suggestions = []

    lower_text = text.lower()

    # --------------------------------------------------------
    # Contact information
    # --------------------------------------------------------

    if not re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    ):

        suggestions.append(
            "Add a professional email address."
        )

    if not re.search(
        r"(?:\+91[\s-]?)?[6-9]\d{9}",
        text
    ):

        suggestions.append(
            "Add a valid phone number."
        )

    # --------------------------------------------------------
    # Sections
    # --------------------------------------------------------

    if "education" not in lower_text:

        suggestions.append(
            "Add a clear Education section."
        )

    if "project" not in lower_text:

        suggestions.append(
            "Add a Projects section with relevant projects."
        )

    if (
        "experience" not in lower_text
        and "internship" not in lower_text
    ):

        suggestions.append(
            "Add internship, training, or experience details."
        )

    if "skill" not in lower_text:

        suggestions.append(
            "Add a clearly labeled Technical Skills section."
        )

    # --------------------------------------------------------
    # Skills
    # --------------------------------------------------------

    if len(skills) < 5:

        suggestions.append(
            "Add more relevant technical skills that match your target role."
        )

    # --------------------------------------------------------
    # Action verbs
    # --------------------------------------------------------

    if len(action_verbs) < 3:

        suggestions.append(
            "Use stronger action verbs such as Developed, Implemented, Designed, Tested, and Optimized."
        )

    # --------------------------------------------------------
    # Quantified results
    # --------------------------------------------------------

    if len(quantified_results) == 0:

        suggestions.append(
            "Add measurable results using numbers, percentages, users, modules, or project metrics."
        )

    # --------------------------------------------------------
    # Missing keywords
    # --------------------------------------------------------

    if missing_keywords:

        suggestions.append(
            "Consider adding relevant missing job skills: "
            + ", ".join(
                missing_keywords
            )
            + "."
        )

    # --------------------------------------------------------
    # Final message
    # --------------------------------------------------------

    if not suggestions:

        suggestions.append(
            "Resume structure and content coverage look good. Continue tailoring it to each job description."
        )

    return suggestions


# ============================================================
# MAIN ATS CALCULATION
# ============================================================

def calculate_ats_score(
    resume_text,
    job_description=""
):

    skills = detect_skills(
        resume_text
    )

    action_verbs = detect_action_verbs(
        resume_text
    )

    quantified_results = detect_quantified_results(
        resume_text
    )

    skills_score = calculate_skill_score(
        skills
    )

    section_score = calculate_section_score(
        resume_text
    )

    contact_score = calculate_contact_score(
        resume_text
    )

    length_score = calculate_length_score(
        resume_text
    )

    action_score = calculate_action_score(
        action_verbs
    )

    quantified_score = calculate_quantified_score(
        quantified_results
    )

    (
        keyword_score,
        keyword_match_percentage,
        matched_keywords,
        missing_keywords
    ) = calculate_keyword_match(
        resume_text,
        job_description
    )

    overall_score = (
        skills_score
        + section_score
        + contact_score
        + length_score
        + action_score
        + quantified_score
        + keyword_score
    )

    overall_score = round(
        min(
            overall_score,
            100
        ),
        2
    )

    suggestions = generate_suggestions(
        resume_text,
        skills,
        action_verbs,
        quantified_results,
        matched_keywords,
        missing_keywords
    )

    return {

        "overall_score": overall_score,

        "skills_score": skills_score,

        "section_score": section_score,

        "contact_score": contact_score,

        "length_score": length_score,

        "action_score": action_score,

        "quantified_score": quantified_score,

        "keyword_score": keyword_score,

        "keyword_match_percentage": keyword_match_percentage,

        "matched_keywords": matched_keywords,

        "missing_keywords": missing_keywords,

        "detected_skills": skills,

        "action_verbs": action_verbs,

        "quantified_results": quantified_results,

        "suggestions": suggestions
    }


# ============================================================
# HOME PAGE
# ============================================================

@app.route("/")
def index():

    return render_template(
        "index.html"
    )


# ============================================================
# REGISTER
# ============================================================

@app.route(
    "/register",
    methods=["GET", "POST"]
)
def register():

    if request.method == "GET":

        return render_template(
            "register.html"
        )

    full_name = request.form.get(
        "full_name",
        ""
    ).strip()

    email = request.form.get(
        "email",
        ""
    ).strip().lower()

    password = request.form.get(
        "password",
        ""
    )

    confirm_password = request.form.get(
        "confirm_password",
        ""
    )

    if (
        not full_name
        or not email
        or not password
        or not confirm_password
    ):

        flash(
            "Please fill in all required fields.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    if not re.match(
        r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$",
        email
    ):

        flash(
            "Please enter a valid email address.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    if password != confirm_password:

        flash(
            "Passwords do not match.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    if len(password) < 6:

        flash(
            "Password must contain at least 6 characters.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    connection = get_db_connection()

    if connection is None:

        flash(
            "Database connection failed. Please make sure MySQL is running.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    cursor = None

    try:

        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT id
            FROM users
            WHERE email = %s
            """,
            (email,)
        )

        existing_user = cursor.fetchone()

        if existing_user:

            flash(
                "Email already registered. Please login.",
                "warning"
            )

            return redirect(
                url_for("login")
            )

        hashed_password = generate_password_hash(
            password
        )

        cursor.execute(
            """
            INSERT INTO users
            (
                full_name,
                email,
                password,
                role
            )
            VALUES (%s, %s, %s, %s)
            """,
            (
                full_name,
                email,
                hashed_password,
                "user"
            )
        )

        connection.commit()

        flash(
            "Account created successfully. Please login.",
            "success"
        )

        return redirect(
            url_for("login")
        )

    except mysql.connector.IntegrityError as e:

        connection.rollback()

        print(
            "Registration integrity error:",
            e
        )

        flash(
            "This email is already registered.",
            "warning"
        )

        return redirect(
            url_for("login")
        )

    except Error as e:

        connection.rollback()

        print(
            "Registration database error:",
            e
        )

        flash(
            "Registration failed. Please try again.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    except Exception as e:

        connection.rollback()

        print(
            "Registration error:",
            e
        )

        flash(
            "An unexpected error occurred during registration.",
            "danger"
        )

        return redirect(
            url_for("register")
        )

    finally:

        if cursor:

            try:
                cursor.close()
            except Exception:
                pass

        try:
            connection.close()
        except Exception:
            pass


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if request.method == "GET":

        return render_template(
            "login.html"
        )

    email = request.form.get(
        "email",
        ""
    ).strip().lower()

    password = request.form.get(
        "password",
        ""
    )

    if not email or not password:

        flash(
            "Please enter email and password.",
            "danger"
        )

        return redirect(
            url_for("login")
        )

    connection = get_db_connection()

    if connection is None:

        flash(
            "Database connection failed. Please make sure MySQL is running.",
            "danger"
        )

        return redirect(
            url_for("login")
        )

    cursor = None

    try:

        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute(
            """
            SELECT
                id,
                full_name,
                email,
                password,
                role
            FROM users
            WHERE email = %s
            LIMIT 1
            """,
            (email,)
        )

        user = cursor.fetchone()

        if user:

            stored_password = user.get(
                "password"
            )

            if (
                stored_password
                and check_password_hash(
                    stored_password,
                    password
                )
            ):

                session.clear()

                session["user_id"] = user["id"]

                session["user_name"] = user["full_name"]

                session["user_email"] = user["email"]

                session["user_role"] = (
                    user["role"]
                    or "user"
                )

                session.permanent = False

                flash(
                    "Login successful.",
                    "success"
                )

                return redirect(
                    url_for("dashboard")
                )

        flash(
            "Invalid email or password.",
            "danger"
        )

        return redirect(
            url_for("login")
        )

    except Error as e:

        print(
            "Login database error:",
            e
        )

        flash(
            "Login failed. Please try again.",
            "danger"
        )

        return redirect(
            url_for("login")
        )

    except Exception as e:

        print(
            "Login error:",
            e
        )

        flash(
            "An unexpected error occurred during login.",
            "danger"
        )

        return redirect(
            url_for("login")
        )

    finally:

        if cursor:

            try:
                cursor.close()
            except Exception:
                pass

        try:
            connection.close()
        except Exception:
            pass


# ============================================================
# DASHBOARD
# ============================================================

@app.route("/dashboard")
def dashboard():

    # --------------------------------------------------------
    # LOGIN REQUIRED
    # --------------------------------------------------------

    if "user_id" not in session:

        flash(
            "Please login to access your dashboard.",
            "warning"
        )

        return redirect(
            url_for("login")
        )

    connection = get_db_connection()

    resumes = []

    if connection:

        cursor = None

        try:

            cursor = connection.cursor(
                dictionary=True
            )

            cursor.execute(
                """
                SELECT
                    id,
                    filename,
                    skills,
                    uploaded_at
                FROM resumes
                WHERE user_id = %s
                ORDER BY uploaded_at DESC
                """,
                (session["user_id"],)
            )

            resumes = cursor.fetchall()

        except Error as e:

            print(
                "Dashboard database error:",
                e
            )

        finally:

            if cursor:

                try:
                    cursor.close()
                except Exception:
                    pass

            try:
                connection.close()
            except Exception:
                pass

    return render_template(
        "dashboard.html",
        resumes=resumes,
        resume_count=len(resumes),
        user_name=session.get(
            "user_name",
            "User"
        )
    )


# ============================================================
# UPLOAD RESUME
# ============================================================

@app.route(
    "/upload-resume",
    methods=["GET", "POST"]
)
def upload_resume():

    # --------------------------------------------------------
    # LOGIN REQUIRED
    # --------------------------------------------------------

    if "user_id" not in session:

        flash(
            "Please login before uploading a resume.",
            "warning"
        )

        return redirect(
            url_for("login")
        )

    # --------------------------------------------------------
    # GET
    # --------------------------------------------------------

    if request.method == "GET":

        return render_template(
            "upload_resume.html"
        )

    # --------------------------------------------------------
    # FORM DATA
    # --------------------------------------------------------

    file = request.files.get(
        "resume"
    )

    job_title = request.form.get(
        "job_title",
        ""
    ).strip()

    company = request.form.get(
        "company",
        ""
    ).strip()

    location = request.form.get(
        "location",
        ""
    ).strip()

    job_description = request.form.get(
        "job_description",
        ""
    ).strip()

    # --------------------------------------------------------
    # FILE VALIDATION
    # --------------------------------------------------------

    if not file or not file.filename:

        flash(
            "Please select a resume file.",
            "danger"
        )

        return redirect(
            url_for("upload_resume")
        )

    if not allowed_file(file.filename):

        flash(
            "Only PDF and DOCX files are allowed.",
            "danger"
        )

        return redirect(
            url_for("upload_resume")
        )

    original_filename = secure_filename(
        file.filename
    )

    if not original_filename:

        flash(
            "Invalid filename.",
            "danger"
        )

        return redirect(
            url_for("upload_resume")
        )

    # --------------------------------------------------------
    # UNIQUE FILE NAME
    # --------------------------------------------------------

    timestamp = datetime.now().strftime(
        "%Y%m%d%H%M%S%f"
    )

    filename = (
        timestamp
        + "_"
        + original_filename
    )

    os.makedirs(
        app.config["UPLOAD_FOLDER"],
        exist_ok=True
    )

    file_path = os.path.join(
        app.config["UPLOAD_FOLDER"],
        filename
    )

    connection = None
    cursor = None

    try:

        # ----------------------------------------------------
        # SAVE FILE
        # ----------------------------------------------------

        file.save(
            file_path
        )

        # ----------------------------------------------------
        # EXTRACT RESUME TEXT
        # ----------------------------------------------------

        resume_text = extract_resume_text(
            file_path
        )

        if not resume_text:

            try:

                if os.path.exists(file_path):

                    os.remove(
                        file_path
                    )

            except OSError:

                pass

            flash(
                "Could not extract text from this resume. Please use a text-based PDF or DOCX file.",
                "danger"
            )

            return redirect(
                url_for("upload_resume")
            )

        # ----------------------------------------------------
        # ATS ANALYSIS
        # ----------------------------------------------------

        ats_data = calculate_ats_score(
            resume_text,
            job_description
        )

        # ----------------------------------------------------
        # DATABASE CONNECTION
        # ----------------------------------------------------

        connection = get_db_connection()

        if connection is None:

            try:

                if os.path.exists(file_path):

                    os.remove(
                        file_path
                    )

            except OSError:

                pass

            flash(
                "Database connection failed.",
                "danger"
            )

            return redirect(
                url_for("upload_resume")
            )

        cursor = connection.cursor()

        # ----------------------------------------------------
        # SAVE RESUME
        # ----------------------------------------------------

        cursor.execute(
            """
            INSERT INTO resumes
            (
                user_id,
                filename,
                extracted_text,
                skills
            )
            VALUES (%s, %s, %s, %s)
            """,
            (
                session["user_id"],
                original_filename,
                resume_text,
                ", ".join(
                    ats_data["detected_skills"]
                )
            )
        )

        resume_id = cursor.lastrowid

        # ----------------------------------------------------
        # SAVE JOB DESCRIPTION
        # ----------------------------------------------------

        job_description_id = None

        if (
            job_title
            or company
            or location
            or job_description
        ):

            cursor.execute(
                """
                INSERT INTO job_descriptions
                (
                    user_id,
                    job_title,
                    company,
                    location,
                    description
                )
                VALUES (%s, %s, %s, %s, %s)
                """,
                (
                    session["user_id"],
                    job_title,
                    company,
                    location,
                    job_description
                )
            )

            job_description_id = (
                cursor.lastrowid
            )

        # ----------------------------------------------------
        # SAVE ANALYSIS RESULT
        # ----------------------------------------------------

        cursor.execute(
            """
            INSERT INTO analysis_results
            (
                resume_id,
                resume_score,
                text_length,
                skills_count,
                detected_skills,
                suggestions
            )
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                resume_id,

                ats_data["overall_score"],

                len(resume_text),

                len(
                    ats_data["detected_skills"]
                ),

                json.dumps(
                    ats_data["detected_skills"]
                ),

                json.dumps(
                    ats_data["suggestions"]
                )
            )
        )

        # ----------------------------------------------------
        # SAVE ATS SCORE
        # ----------------------------------------------------

        cursor.execute(
            """
            INSERT INTO ats_scores
            (
                resume_id,
                job_description_id,
                overall_score,
                skills_score,
                section_score,
                contact_score,
                length_score,
                action_score,
                quantified_score,
                keyword_score,
                keyword_match_percentage,
                matched_keywords,
                missing_keywords,
                detected_skills,
                action_verbs,
                quantified_results,
                suggestions
            )
            VALUES (
                %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s
            )
            """,
            (
                resume_id,

                job_description_id,

                ats_data["overall_score"],

                ats_data["skills_score"],

                ats_data["section_score"],

                ats_data["contact_score"],

                ats_data["length_score"],

                ats_data["action_score"],

                ats_data["quantified_score"],

                ats_data["keyword_score"],

                ats_data[
                    "keyword_match_percentage"
                ],

                json.dumps(
                    ats_data[
                        "matched_keywords"
                    ]
                ),

                json.dumps(
                    ats_data[
                        "missing_keywords"
                    ]
                ),

                json.dumps(
                    ats_data[
                        "detected_skills"
                    ]
                ),

                json.dumps(
                    ats_data[
                        "action_verbs"
                    ]
                ),

                json.dumps(
                    ats_data[
                        "quantified_results"
                    ]
                ),

                json.dumps(
                    ats_data[
                        "suggestions"
                    ]
                )
            )
        )

        # ----------------------------------------------------
        # COMMIT
        # ----------------------------------------------------

        connection.commit()

        # ----------------------------------------------------
        # CLOSE DATABASE
        # ----------------------------------------------------

        cursor.close()
        cursor = None

        connection.close()
        connection = None

        # ----------------------------------------------------
        # SHOW ANALYSIS PAGE
        # ----------------------------------------------------

        return render_template(
            "resume_analysis.html",

            filename=original_filename,

            resume_text=resume_text,

            resume_score=ats_data[
                "overall_score"
            ],

            skills=ats_data[
                "detected_skills"
            ],

            action_verbs=ats_data[
                "action_verbs"
            ],

            quantified_results=ats_data[
                "quantified_results"
            ],

            matched_keywords=ats_data[
                "matched_keywords"
            ],

            missing_keywords=ats_data[
                "missing_keywords"
            ],

            job_title=job_title,

            company=company,

            location=location,

            job_description=job_description,

            ats_data=ats_data
        )

    except Exception as e:

        print(
            "Resume upload/analysis error:",
            e
        )

        if connection:

            try:
                connection.rollback()
            except Exception:
                pass

        try:

            if os.path.exists(file_path):

                os.remove(
                    file_path
                )

        except OSError:

            pass

        flash(
            "An error occurred while processing the resume.",
            "danger"
        )

        return redirect(
            url_for("upload_resume")
        )

    finally:

        if cursor:

            try:
                cursor.close()
            except Exception:
                pass

        if connection:

            try:
                connection.close()
            except Exception:
                pass


# ============================================================
# ANALYZE SAVED RESUME
# ============================================================

@app.route(
    "/analyze-resume",
    methods=["GET", "POST"]
)
def analyze_resume():

    # --------------------------------------------------------
    # LOGIN REQUIRED
    # --------------------------------------------------------

    if "user_id" not in session:

        flash(
            "Please login before analyzing a resume.",
            "warning"
        )

        return redirect(
            url_for("login")
        )

    # ========================================================
    # GET REQUEST
    #
    # This is used when clicking the Analyze button
    # from the Dashboard.
    #
    # Example:
    #
    # /analyze-resume?resume_id=5
    #
    # ========================================================

    if request.method == "GET":

        resume_id = request.args.get(
            "resume_id",
            type=int
        )

        if not resume_id:

            flash(
                "Please select a resume to analyze.",
                "warning"
            )

            return redirect(
                url_for("dashboard")
            )

        connection = get_db_connection()

        if connection is None:

            flash(
                "Database connection failed.",
                "danger"
            )

            return redirect(
                url_for("dashboard")
            )

        cursor = None

        try:

            cursor = connection.cursor(
                dictionary=True
            )

            # ------------------------------------------------
            # IMPORTANT SECURITY CHECK
            #
            # The resume must belong to the logged-in user.
            # ------------------------------------------------

            cursor.execute(
                """
                SELECT
                    id,
                    user_id,
                    filename,
                    extracted_text,
                    skills,
                    uploaded_at
                FROM resumes
                WHERE id = %s
                AND user_id = %s
                LIMIT 1
                """,
                (
                    resume_id,
                    session["user_id"]
                )
            )

            resume = cursor.fetchone()

            if not resume:

                flash(
                    "Resume not found.",
                    "danger"
                )

                return redirect(
                    url_for("dashboard")
                )

            resume_text = resume.get(
                "extracted_text"
            ) or ""

            if not resume_text:

                flash(
                    "No extracted text is available for this resume.",
                    "danger"
                )

                return redirect(
                    url_for("dashboard")
                )

            # ------------------------------------------------
            # AUTOMATIC GENERAL ATS ANALYSIS
            #
            # No job description is required when analyzing
            # directly from the Dashboard.
            # ------------------------------------------------

            ats_data = calculate_ats_score(
                resume_text,
                ""
            )

            # ------------------------------------------------
            # SAVE NEW ANALYSIS RESULT
            # ------------------------------------------------

            cursor.execute(
                """
                INSERT INTO analysis_results
                (
                    resume_id,
                    resume_score,
                    text_length,
                    skills_count,
                    detected_skills,
                    suggestions
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    resume_id,

                    ats_data["overall_score"],

                    len(resume_text),

                    len(
                        ats_data[
                            "detected_skills"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "detected_skills"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "suggestions"
                        ]
                    )
                )
            )

            # ------------------------------------------------
            # SAVE ATS SCORE
            #
            # This is a general resume analysis, so there is
            # no job description ID.
            # ------------------------------------------------

            cursor.execute(
                """
                INSERT INTO ats_scores
                (
                    resume_id,
                    job_description_id,
                    overall_score,
                    skills_score,
                    section_score,
                    contact_score,
                    length_score,
                    action_score,
                    quantified_score,
                    keyword_score,
                    keyword_match_percentage,
                    matched_keywords,
                    missing_keywords,
                    detected_skills,
                    action_verbs,
                    quantified_results,
                    suggestions
                )
                VALUES (
                    %s, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s
                )
                """,
                (
                    resume_id,

                    None,

                    ats_data["overall_score"],

                    ats_data["skills_score"],

                    ats_data["section_score"],

                    ats_data["contact_score"],

                    ats_data["length_score"],

                    ats_data["action_score"],

                    ats_data["quantified_score"],

                    ats_data["keyword_score"],

                    ats_data[
                        "keyword_match_percentage"
                    ],

                    json.dumps(
                        ats_data[
                            "matched_keywords"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "missing_keywords"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "detected_skills"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "action_verbs"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "quantified_results"
                        ]
                    ),

                    json.dumps(
                        ats_data[
                            "suggestions"
                        ]
                    )
                )
            )

            connection.commit()

            return render_template(
                "resume_analysis.html",

                filename=resume[
                    "filename"
                ],

                resume_text=resume_text,

                resume_score=ats_data[
                    "overall_score"
                ],

                skills=ats_data[
                    "detected_skills"
                ],

                action_verbs=ats_data[
                    "action_verbs"
                ],

                quantified_results=ats_data[
                    "quantified_results"
                ],

                matched_keywords=ats_data[
                    "matched_keywords"
                ],

                missing_keywords=ats_data[
                    "missing_keywords"
                ],

                job_title="",

                company="",

                location="",

                job_description="",

                ats_data=ats_data
            )

        except Error as e:

            if connection:

                try:
                    connection.rollback()
                except Exception:
                    pass

            print(
                "Saved resume analysis database error:",
                e
            )

            flash(
                "Could not analyze the selected resume.",
                "danger"
            )

            return redirect(
                url_for("dashboard")
            )

        except Exception as e:

            if connection:

                try:
                    connection.rollback()
                except Exception:
                    pass

            print(
                "Saved resume analysis error:",
                e
            )

            flash(
                "An error occurred while analyzing the resume.",
                "danger"
            )

            return redirect(
                url_for("dashboard")
            )

        finally:

            if cursor:

                try:
                    cursor.close()
                except Exception:
                    pass

            try:
                connection.close()
            except Exception:
                pass

    # ========================================================
    # POST REQUEST
    #
    # This preserves the existing upload/analyze form flow.
    # ========================================================

    return upload_resume()


# ============================================================
# JOBS
# ============================================================

@app.route("/jobs")
def jobs():

    # --------------------------------------------------------
    # LOGIN REQUIRED
    # --------------------------------------------------------

    if "user_id" not in session:

        flash(
            "Please login to view jobs.",
            "warning"
        )

        return redirect(
            url_for("login")
        )

    connection = get_db_connection()

    jobs_list = []

    if connection:

        cursor = None

        try:

            cursor = connection.cursor(
                dictionary=True
            )

            # ------------------------------------------------
            # Read existing jobs table
            #
            # IMPORTANT:
            # This table is NOT created or modified here.
            # ------------------------------------------------

            cursor.execute(
                """
                SELECT
                    id,
                    title,
                    company,
                    location,
                    skills,
                    description
                FROM jobs
                ORDER BY id DESC
                """
            )

            jobs_list = cursor.fetchall()

        except Error as e:

            print(
                "Jobs database error:",
                e
            )

            jobs_list = []

        finally:

            if cursor:

                try:
                    cursor.close()
                except Exception:
                    pass

            try:
                connection.close()
            except Exception:
                pass

    return render_template(
        "jobs.html",
        jobs=jobs_list
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    session.clear()

    flash(
        "You have been logged out.",
        "success"
    )

    return redirect(
        url_for("index")
    )


# ============================================================
# 404 ERROR HANDLER
# ============================================================

@app.errorhandler(404)
def page_not_found(error):

    return render_template(
        "404.html"
    ), 404


# ============================================================
# 413 ERROR HANDLER
# ============================================================

@app.errorhandler(413)
def file_too_large(error):

    flash(
        "File is too large. Maximum size is 10 MB.",
        "danger"
    )

    if "user_id" in session:

        return redirect(
            url_for("upload_resume")
        )

    return redirect(
        url_for("login")
    )


# ============================================================
# APPLICATION START
# ============================================================

if __name__ == "__main__":

    # --------------------------------------------------------
    # Create upload folder
    # --------------------------------------------------------

    os.makedirs(
        app.config["UPLOAD_FOLDER"],
        exist_ok=True
    )

    # --------------------------------------------------------
    # Initialize database
    # --------------------------------------------------------

    database_ready = initialize_database()

    if not database_ready:

        print(
            "\nWARNING: Database initialization failed."
        )

        print(
            "Make sure MySQL is running and DB_CONFIG is correct.\n"
        )

    # --------------------------------------------------------
    # SERVER INFORMATION
    # --------------------------------------------------------

    print(
        "=============================================="
    )

    print(
        "AI Resume Screening System"
    )

    print(
        "=============================================="
    )

    print(
        "Homepage : http://127.0.0.1:5000/"
    )

    print(
        "Login    : http://127.0.0.1:5000/login"
    )

    print(
        "Register : http://127.0.0.1:5000/register"
    )

    print(
        "Dashboard: http://127.0.0.1:5000/dashboard"
    )

    print(
        "Upload   : http://127.0.0.1:5000/upload-resume"
    )

    print(
        "Jobs     : http://127.0.0.1:5000/jobs"
    )

    print(
        "=============================================="
    )

    # --------------------------------------------------------
    # START FLASK
    # --------------------------------------------------------

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )